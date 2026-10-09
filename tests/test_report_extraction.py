import io
import json
import os

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

import main
from docx_generator import docx_generator
from transcript_to_json import report_extractor
from transcript_to_json.report_schema import REPORT_FIELDS, empty_report

TRANSCRIPT = "Caller: Machine RP-204 displays E204. Restarting did not help."
MODEL_FIELDS = {
    "machine_number": "RP-204",
    "customer": None,
    "problem": "Machine displays E204.",
    "alarm_code": "E204",
    "actions_taken": "Restarting did not help.",
}


@pytest.fixture
def model_response(monkeypatch):
    monkeypatch.setenv("REPORT_MODEL", "qwen3:4b")
    monkeypatch.delenv("REPORT_DEVICE", raising=False)
    body = {"done": True, "message": {"content": json.dumps(MODEL_FIELDS)}}
    original_client = httpx.Client

    def respond(request):
        payload = json.loads(request.content)
        assert request.url.path == "/api/chat"
        assert payload["model"] == "qwen3:4b"
        assert payload["think"] is False
        assert payload["options"]["temperature"] == 0.7
        assert payload["options"]["top_p"] == 0.8
        assert payload["options"]["top_k"] == 20
        assert payload["options"]["min_p"] == 0
        assert payload["options"]["num_predict"] == -1
        assert all(value is None for value in request.extensions["timeout"].values())
        assert payload["messages"][-1]["content"] == TRANSCRIPT
        return httpx.Response(200, json=body)

    monkeypatch.setattr(report_extractor.httpx, "Client", lambda **kwargs: original_client(
        transport=httpx.MockTransport(respond), **kwargs,
    ))
    return body


def test_transcript_populates_report_fields(model_response):
    response = TestClient(main.app).post("/extract-report", json={"transcript": TRANSCRIPT})
    assert response.status_code == 200
    result = response.json()
    assert result["mode"] == "llm"
    assert result["warning"] is None
    assert set(result["fields"]) == set(REPORT_FIELDS)
    assert result["fields"]["[machinenummer]"] == "RP-204"
    assert result["fields"]["[probleem]"] == "Machine displays E204."
    assert result["fields"]["[alarmcode of exacte tekst]"] == "E204"
    assert result["fields"]["[transcriptie]"] == TRANSCRIPT
    assert result["fields"]["[klant]"] == ""
    assert result["fields"]["[goedgekeurd]"] == ""


@pytest.mark.parametrize("body", [
    [], {"done": False}, {"done": True, "done_reason": "length"},
    {"done": True}, {"done": True, "message": {"content": "not JSON"}},
])
def test_bad_model_output_uses_generic_manual_fallback(model_response, body):
    model_response.clear()
    # A response without usable JSON should always reach the same fallback.
    model_response.update(body if isinstance(body, dict) else {"message": body})
    result = TestClient(main.app).post("/extract-report", json={"transcript": TRANSCRIPT}).json()
    assert result["mode"] == "manual"
    assert result["fields"]["[transcriptie]"] == TRANSCRIPT
    assert result["warning"] == "Automatic filling is unavailable. Complete the remaining details manually."


@pytest.mark.parametrize("error", [
    httpx.ConnectError("Ollama unavailable"),
    httpx.ReadTimeout("Model timed out"),
    ValueError("Invalid result"),
    RuntimeError("Unexpected failure"),
])
def test_extraction_errors_preserve_manual_entry(monkeypatch, error):
    def fail(transcript, model=None):
        raise error
    monkeypatch.setattr(main, "extract_call_fields_from_transcript", fail)
    response = TestClient(main.app).post("/extract-report", json={"transcript": TRANSCRIPT})
    assert response.status_code == 200
    result = response.json()
    assert result["mode"] == "manual"
    assert result["fields"]["[transcriptie]"] == TRANSCRIPT
    assert all(value == "" for key, value in result["fields"].items() if key != "[transcriptie]")
    assert result["warning"] == "Automatic filling is unavailable. Complete the remaining details manually."


@pytest.mark.parametrize("payload", [{}, {"transcript": ""}, {"transcript": "   "}])
def test_blank_requests_do_not_call_model(monkeypatch, payload):
    monkeypatch.setattr(main, "extract_call_fields_from_transcript", lambda text, model=None: pytest.fail("Model called"))
    assert TestClient(main.app).post("/extract-report", json=payload).status_code == 422


@pytest.mark.parametrize("model", ["qwen3:8b", "qwen3:1.7b"])
def test_model_choice_and_cpu_setting(monkeypatch, model):
    monkeypatch.setenv("REPORT_MODEL", "qwen3:4b")
    monkeypatch.setenv("REPORT_DEVICE", "cpu")
    original_client = httpx.Client

    def respond(request):
        payload = json.loads(request.content)
        assert payload["model"] == model
        assert payload["options"]["num_gpu"] == 0
        return httpx.Response(200, json={"done": True, "message": {"content": json.dumps(MODEL_FIELDS)}})

    monkeypatch.setattr(report_extractor.httpx, "Client", lambda **kwargs: original_client(
        transport=httpx.MockTransport(respond), **kwargs,
    ))
    response = TestClient(main.app).post("/extract-report", json={"transcript": TRANSCRIPT, "model": model})
    assert response.status_code == 200
    assert response.json()["fields"]["[alarmcode of exacte tekst]"] == "E204"
    assert os.environ["REPORT_MODEL"] == "qwen3:4b"


def test_disabled_model_preserves_full_transcript(monkeypatch):
    monkeypatch.setenv("REPORT_MODEL", "qwen3:4b")
    monkeypatch.setattr(report_extractor.httpx, "Client", lambda **kwargs: pytest.fail("Model contacted"))
    result = TestClient(main.app).post("/extract-report", json={"transcript": TRANSCRIPT, "model": "none"}).json()
    assert result["mode"] == "manual"
    assert result["fields"]["[transcriptie]"] == TRANSCRIPT


@pytest.mark.parametrize("installed, configured, expected", [
    (["qwen3:1.7b", "qwen3:4b"], "qwen3:4b", "qwen3:4b"),
    (["qwen3:1.7b", "qwen3:4b"], "qwen3:1.7b", "qwen3:1.7b"),
    (["qwen3:1.7b"], "qwen3:4b", "qwen3:1.7b"),
    (["qwen3:4b"], "none", "none"),
    ([], "qwen3:4b", "none"),
])
def test_selector_lists_installed_models_and_respects_default(monkeypatch, installed, configured, expected):
    monkeypatch.setenv("REPORT_MODEL", configured)
    original_client = httpx.Client

    def respond(request):
        assert request.url.path == "/api/tags"
        return httpx.Response(200, json={"models": [{"name": name} for name in installed]})

    monkeypatch.setattr(main.httpx, "Client", lambda **kwargs: original_client(
        transport=httpx.MockTransport(respond), **kwargs,
    ))
    response = TestClient(main.app).get("/report-models")
    assert response.status_code == 200
    assert response.json()["default"] == expected
    assert {model["name"] for model in response.json()["models"] if model["installed"]} == set(installed)


def test_model_selector_remains_usable_without_ollama(monkeypatch):
    original_client = httpx.Client

    def fail(request):
        raise httpx.ConnectError("Ollama unavailable")

    monkeypatch.setattr(main.httpx, "Client", lambda **kwargs: original_client(
        transport=httpx.MockTransport(fail), **kwargs,
    ))
    result = TestClient(main.app).get("/report-models").json()
    assert result["default"] == "none"
    assert all(not model["installed"] for model in result["models"])


@pytest.mark.parametrize("dialogue", ["Checking the machine.", "正在检查机器。"])
def test_long_calls_extract_all_sections_and_combine_details(monkeypatch, dialogue):
    monkeypatch.setenv("REPORT_MODEL", "qwen3:4b")
    transcript = TRANSCRIPT + "\n" + "".join(
        f"Harm: Check {index}. Machine RP-204. {dialogue}\n" for index in range(6000)
    ) + "Caller: We replaced the sensor. Our company is Example Packaging."
    assert len(transcript) > 200000
    prepared = report_extractor.prepare_transcript(transcript)
    requests = []
    original_client = httpx.Client

    def respond(request):
        payload = json.loads(request.content)
        section = payload["messages"][-1]["content"]
        requests.append(payload)
        fields = {"machine_number": None, "customer": None, "problem": None,
                  "alarm_code": None, "actions_taken": None}
        facts = {
            "machine_number": ("RP-204", "Machine RP-204"),
            "alarm_code": ("E204", "displays E204"),
            "customer": ("Example Packaging", "Our company is Example Packaging"),
            "actions_taken": ("Restarting did not help.", "Restarting did not help."),
        }
        if "We replaced the sensor." in section:
            facts["actions_taken"] = ("We replaced the sensor.", "We replaced the sensor.")
        for name, (value, text) in facts.items():
            if text in section:
                fields[name] = value
        return httpx.Response(200, json={"done": True, "message": {"content": json.dumps(fields)}})

    monkeypatch.setattr(report_extractor.httpx, "Client", lambda **kwargs: original_client(
        transport=httpx.MockTransport(respond), **kwargs,
    ))
    response = TestClient(main.app).post("/extract-report", json={"transcript": transcript})
    assert response.status_code == 200
    result = response.json()
    assert result["mode"] == "llm"
    assert result["fields"]["[machinenummer]"] == "RP-204"
    assert result["fields"]["[alarmcode of exacte tekst]"] == "E204"
    assert result["fields"]["[klant]"] == "Example Packaging"
    assert result["fields"]["[acties en resultaten]"] == "Restarting did not help.\nWe replaced the sensor."
    assert result["fields"]["[transcriptie]"] == transcript
    assert len(requests) > 1

    covered_until = 0
    for payload in requests:
        section = payload["messages"][-1]["content"]
        start = prepared.index(section)
        assert start <= covered_until  # No transcript gaps between model requests.
        assert len(section.encode("utf-8")) <= 6000
        assert payload["think"] is False
        assert payload["options"]["num_predict"] == -1
        covered_until = start + len(section)
    assert covered_until == len(prepared)
    assert all(payload["keep_alive"] == "5m" for payload in requests[:-1])
    assert requests[-1]["keep_alive"] == 0


def test_section_overlap_keeps_boundary_dialogue_together():
    dialogue = "Harm: Which company?\nCaller: Example Packaging."
    transcript = "x" * 5980 + "\n" + dialogue
    sections = list(report_extractor.split_transcript(transcript))
    assert len(sections) == 2
    assert any(dialogue in section for section in sections)


def test_unknown_values_and_speaker_labels_stay_empty(model_response):
    values = dict(MODEL_FIELDS, customer="not mentioned", contact_person="Caller", symptoms="   ")
    model_response["message"]["content"] = json.dumps(values)
    result = report_extractor.extract_call_fields_from_transcript(TRANSCRIPT)
    assert result.customer is None
    assert result.contact_person is None
    assert result.symptoms is None


def test_timestamp_fragments_are_joined_without_merging_different_speakers(monkeypatch):
    transcript = (
        "[00:00:01] Harm: The problem is that the sealing\n"
        "[00:00:05] Harm: station jams when we increase the speed.\n"
        "[00:00:09] Caller: Our machine is RP-315.\n"
        "[00:00:12] Harm: Has it been restarted?"
    )
    expected = (
        "Harm: The problem is that the sealing station jams when we increase the speed.\n"
        "Caller: Our machine is RP-315.\n"
        "Harm: Has it been restarted?"
    )
    assert report_extractor.prepare_transcript(transcript) == expected
    original_client = httpx.Client

    def respond(request):
        payload = json.loads(request.content)
        assert payload["messages"][-1]["content"] == expected
        fields = dict(MODEL_FIELDS, machine_number="RP-315", alarm_code=None,
                      problem="Sealing station jams when speed increases.", actions_taken=None)
        return httpx.Response(200, json={"done": True, "message": {"content": json.dumps(fields)}})

    monkeypatch.setattr(report_extractor.httpx, "Client", lambda **kwargs: original_client(
        transport=httpx.MockTransport(respond), **kwargs,
    ))
    response = TestClient(main.app).post("/extract-report", json={"transcript": transcript})
    result = response.json()
    assert result["mode"] == "llm"
    assert result["fields"]["[probleem]"] == "Sealing station jams when speed increases."
    assert result["fields"]["[transcriptie]"] == transcript


def test_fields_cover_word_template():
    assert set(REPORT_FIELDS) == docx_generator.find_unreplaced_placeholders(Document(docx_generator.template_docx))


def test_manual_values_reach_downloaded_docx(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_generator, "script_directory", tmp_path)
    fields = empty_report()
    fields.update({"[klant]": "Manual customer", "[transcriptie]": TRANSCRIPT})
    response = TestClient(main.app).post("/export", json={"transcript": json.dumps(fields)})
    assert response.status_code == 200
    document = Document(io.BytesIO(response.content))
    text = " ".join(node.text or "" for node in document.element.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
    assert "Manual customer" in text
    assert TRANSCRIPT in text
    assert docx_generator.find_unreplaced_placeholders(document) == set()
