import json

from docx import Document
from docx.oxml.ns import qn

from docx_generator.docx_generator import find_unreplaced_placeholders, replace_text_in_document


def test_multiline_transcript_has_word_line_breaks(tmp_path):
    document = Document()
    document.add_paragraph("Transcript: [transcriptie]")
    transcript = "First line\r\nSecond line\nThird line"
    replacements = {"[transcriptie]": transcript}

    replace_text_in_document(document, json.dumps(replacements))
    output_path = tmp_path / "transcript.docx"
    document.save(output_path)

    exported_document = Document(output_path)
    paragraph = exported_document.paragraphs[0]
    assert paragraph.text == "Transcript: First line\nSecond line\nThird line"
    assert len(list(paragraph._p.iter(qn("w:br")))) == 2


def test_placeholder_split_across_runs_is_replaced_in_saved_document(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Before [trans")
    paragraph.add_run("criptie] after")
    replacements = {"[transcriptie]": "First line\nSecond line"}

    replace_text_in_document(document, json.dumps(replacements))
    output_path = tmp_path / "split-placeholder.docx"
    document.save(output_path)

    exported_document = Document(output_path)
    exported_paragraph = exported_document.paragraphs[0]
    assert exported_paragraph.text == "Before First line\nSecond line after"
    assert len(list(exported_paragraph._p.iter(qn("w:br")))) == 1
    assert find_unreplaced_placeholders(exported_document) == set()
    text_nodes = list(exported_paragraph._p.iter(qn("w:t")))
    assert text_nodes[-1].text == " after"
    assert text_nodes[-1].get(qn("xml:space")) == "preserve"


def test_empty_replacement_key_keeps_per_node_behavior(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("A")
    paragraph.add_run("B")

    replace_text_in_document(document, json.dumps({"": "-"}))
    output_path = tmp_path / "empty-key.docx"
    document.save(output_path)

    exported_document = Document(output_path)
    assert exported_document.paragraphs[0].text == "-A--B-"
