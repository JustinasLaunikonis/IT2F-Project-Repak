import json
import re
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from docx_generator import docx_generator


def get_document_text(document):
    parts = [document.element]
    for section in document.sections:
        parts.append(section.header._element)
        parts.append(section.first_page_header._element)
        parts.append(section.even_page_header._element)
        parts.append(section.footer._element)
        parts.append(section.first_page_footer._element)
        parts.append(section.even_page_footer._element)

    paragraphs = []
    for part in parts:
        for paragraph in part.iter(qn("w:p")):
            paragraph_text = ""
            for text_node in paragraph.iter(qn("w:t")):
                if text_node.text is not None:
                    paragraph_text += text_node.text
            paragraphs.append(paragraph_text)

    return "\n".join(paragraphs)


def assert_export_contains_values_without_placeholders(output_path, supplied_values):
    saved_document = Document(output_path)
    text = get_document_text(saved_document)

    # Inspect the saved file independently of the generator's placeholder checker.
    remaining_placeholder = re.search(r"\[[^\[\]\r\n]+\]", text)
    if remaining_placeholder is not None:
        raise AssertionError("An unreplaced placeholder remains in the saved document")
    for field, value in supplied_values.items():
        if str(value) not in text:
            raise AssertionError("Missing supplied value for " + field)


def test_example_json_is_present_in_saved_docx(tmp_path, monkeypatch):
    repository_path = Path(__file__).resolve().parent.parent
    example_path = repository_path / "docx_generator" / "template" / "text-input-example.json"
    with example_path.open("r", encoding="utf-8") as example_file:
        example_values = json.load(example_file)

    # Unique synthetic values show whether each example field was inserted.
    supplied_values = {}
    field_number = 1
    for field in example_values:
        marker = "SYNTHETIC_FIELD_" + str(field_number).zfill(3) + "_END"
        supplied_values[field] = marker
        field_number += 1

    input_path = tmp_path / "input.json"
    input_path.write_text(json.dumps(supplied_values), encoding="utf-8")

    monkeypatch.setattr(docx_generator, "script_directory", tmp_path)
    output_path = docx_generator.convert_json_to_docx(input_path)

    assert output_path.parent == tmp_path / "outputs"
    assert output_path.is_file()
    assert_export_contains_values_without_placeholders(output_path, supplied_values)


def test_saved_document_check_rejects_unfilled_placeholder(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Unfilled [trans")
    paragraph.add_run("criptie]")
    output_path = tmp_path / "unfilled.docx"
    document.save(output_path)

    with pytest.raises(AssertionError):
        assert_export_contains_values_without_placeholders(output_path, {})


def test_saved_document_check_rejects_missing_supplied_text(tmp_path):
    document = Document()
    document.add_paragraph("Generated report")
    output_path = tmp_path / "missing-text.docx"
    document.save(output_path)

    with pytest.raises(AssertionError, match="Missing supplied value for"):
        assert_export_contains_values_without_placeholders(
            output_path, {"[transcriptie]": "Synthetic transcript value"}
        )
