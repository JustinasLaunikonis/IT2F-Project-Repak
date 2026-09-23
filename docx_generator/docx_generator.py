import re
from pathlib import Path
from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
import shutil
import os
import subprocess
import platform
from datetime import datetime

# Allows program/directories to be run/accessed from root
script_directory = Path(__file__).resolve().parent
template_docx = script_directory / "template" / "meldingsformulier-template.docx"

# using XML to find the text nodes in the Document is the best approach as XML just returns all the text nodes. This
#   is better than going through paragraphs, tables, footers etc. separately
def print_unreplaced_placeholders(document):

    unreplaced_placeholders = find_unreplaced_placeholders(document)

    if unreplaced_placeholders:
        print("Info: The following placeholders were not replaced in the word document:")
        for unreplaced_placeholder in unreplaced_placeholders:
            print("-", unreplaced_placeholder)
    else:
        print("All placeholders were replaced")

def find_unreplaced_placeholders(document):
    # a set has deduplicated, unordered, unindexed values
    placeholders = set()

    # creates python list containing the document's XML format
    xml_elements = [document.element]

    for section in document.sections:
        xml_elements.extend([
            section.header._element,
            section.first_page_header._element,
            section.even_page_header._element,
            section.footer._element,
            section.first_page_footer._element,
            section.even_page_footer._element
        ])

    for xml_element in xml_elements:
        for paragraph in xml_element.iter(qn("w:p")):
            #turn XML into text. Join text so that a run ending between a placeholder doesn't mess up detection
            text = ""
            for text_node in paragraph.iter(qn("w:t")):
                if text_node.text is not None:
                    text += text_node.text
            # Only text in brackets "[]" is a placeholder in the document, so anything that still has square brackets
            # is a placeholder that hasn't been filled in
            placeholders.update(re.findall(r"\[[^][\r\n]+]", text))

    return placeholders

def replace_text_in_xml(xml_element, replacement_data):
    for text_node in xml_element.iter(qn("w:t")):
        if text_node.text is None:
            continue

        replacement_text = text_node.text
        for key, value in replacement_data.items():
            replacement_text = replacement_text.replace(key, str(value))

        replacement_text = replacement_text.replace("\r\n", "\n")
        replacement_text = replacement_text.replace("\r", "\n")
        lines = replacement_text.split("\n")
        text_node.text = lines[0]

        parent = text_node.getparent()
        insert_at = parent.index(text_node) + 1
        for line in lines[1:]:
            break_node = OxmlElement("w:br")
            parent.insert(insert_at, break_node)
            insert_at += 1
            next_text_node = OxmlElement("w:t")
            next_text_node.text = line
            parent.insert(insert_at, next_text_node)
            insert_at += 1

# Handles both .json files and json strings
def replace_text_in_document(document: DocumentType, text_input : Path | str):

    if isinstance(text_input, Path):
        # Open file and process to json
        with open(text_input, "r", encoding="utf-8") as text_input_json_file:
            text_input_json_file_data = json.load(text_input_json_file)

    elif isinstance(text_input, str):
        # Process String to JSON
        text_input_json_file_data = json.loads(str(text_input))

    else:
        raise TypeError("text_input must be a Path or a str")


    # Main document text
    replace_text_in_xml(document.element, text_input_json_file_data)

    # Headers and Footers
    for section in document.sections:

        replace_text_in_xml(
            section.header._element,
            text_input_json_file_data
        )

        replace_text_in_xml(
            section.first_page_header._element,
            text_input_json_file_data
        )

        replace_text_in_xml(
            section.even_page_header._element,
            text_input_json_file_data
        )

        replace_text_in_xml(
            section.footer._element,
            text_input_json_file_data
        )

        replace_text_in_xml(
            section.first_page_footer._element,
            text_input_json_file_data
        )

        replace_text_in_xml(
            section.even_page_footer._element,
            text_input_json_file_data
        )

#target system is Windows, but this makes testing easier and the program a bit more OS-agnostic
def open_document(file_path):
    file_path = os.path.abspath(file_path)

    if platform.system() == "Windows":
        os.startfile(file_path)

    elif platform.system() == "Linux":
        subprocess.Popen(["xdg-open", file_path])

    elif platform.system() == "Darwin": #macOS
        subprocess.Popen(["open", file_path])

    else:
        print("OS was not recognized. Document could not be opened.")


def convert_json_to_docx(text_input):
    # create folder if missing, does nothing if already exists
    (script_directory / "outputs").mkdir(exist_ok=True)

    # append date and time to working_docx_destination for unique and timestamped outputs
    working_docx_destination = script_directory / "outputs" / (
        Path("output - " + datetime.now().strftime("%d-%m-%Y, %H-%M-%S") + ".docx"))

    iterator = 2

    while Path.exists(working_docx_destination):
        working_docx_destination = script_directory / "outputs" / (
        Path("output - " + datetime.now().strftime("%d-%m-%Y, %H-%M-%S") + "_" + str(iterator) + ".docx"))
        iterator += 1

    # make copy of template file first
    shutil.copyfile(template_docx, working_docx_destination)

    # load document
    document = Document(str(working_docx_destination))

    # Replace the placeholders with JSON data
    replace_text_in_document(document, text_input)

    #Then save
    document.save(str(working_docx_destination))

    # Print the placeholders that were not replaced for debug
    print_unreplaced_placeholders(document)

    # return the file itself
    return Path(working_docx_destination)
