import pytest
from pathlib import Path
from docx_generator import convert_json_to_docx, UnreplacedPlaceholdersError, print_unreplaced_placeholders
from docx import Document
from docx.oxml.ns import qn
import json


def get_xml_text(xml_element):
    text = ""
    for text_node in xml_element.iter(qn("w:t")):
        if text_node.text is None:
            continue

        text += text_node.text

    return str(text)

def get_all_document_text(working_docx_destination):
    document = Document(working_docx_destination)
    text = get_xml_text(document.element)

    for section in document.sections:
        text += "\n" + "\n".join([
            get_xml_text(section.header._element),
            get_xml_text(section.first_page_header._element),
            get_xml_text(section.even_page_header._element),
            get_xml_text(section.footer._element),
            get_xml_text(section.first_page_footer._element),
            get_xml_text(section.even_page_footer._element)
        ])

    print(text)
    return text

# While this test already runs in the docx_generator.py, there it's run from memory. Here the file is opened again.
# I chose not to duplicate the tests and reuse existing tests in docx_generator.py for simplicity
def check_remaining_placeholders(working_docx_destination):
    document = Document(working_docx_destination)
    print_unreplaced_placeholders(document)

def check_supplied_values_appear(working_docx_destination, json_input : Path | str):
    if isinstance(json_input, Path):
        # Open file and process to json
        with open(json_input, "r", encoding="utf-8") as text_input_json_file:
            text_input_json_file_data = json.load(text_input_json_file)

    elif isinstance(json_input, str):
        # Process String to JSON
        text_input_json_file_data = json.loads(str(json_input))

    else:
        raise TypeError("json_input must be a Path or a str")

    text = get_all_document_text(working_docx_destination)

    for placeholder, value in text_input_json_file_data.items():
        assert str(value) in text, ("Missing value. " + placeholder + " : " + value)


def test_generator_json_file_input_does_not_throw_errors():
    script_directory = Path(__file__).resolve().parent

    json_file_input = script_directory / "template" / "text-input-example.json"

    document_path = convert_json_to_docx(json_file_input)
    check_supplied_values_appear(document_path, json_file_input)
    check_remaining_placeholders(document_path)

def test_generator_json_string_input_does_not_throw_errors():

    json_string_input = """{
      "[machinenummer]": "MCH-2024-00158",
      "[distributeur]": "Distributor B.V.",
      "[naam monteur]": "Joost Klein",
      "[contactpersoon]": "Peter Brown",
      "[klant]": "Customer Ltd.",
      "[tijdverschil]": "+1",

      "[datum]": "16-09-2026",
      "[engineer]": "John Smith",
      "[goedgekeurd]": "Smith John",
      "[datum / tijd]": "16-09-2026 / 14:30",

      "[telefoonnummer en/of e-mailadres]": "peter.brown@example.com / +31 6 12345678",
      "[taal]": "English",
      "[naam van Repak medewerker die melding aangenomen heeft]": "John Smith",
      "[machine staat stil / productie beperkt / machine in productie]": "machine is down",
      "[audio link]": "https://audio.website.com/9f33nt9gievnrk",
      "[betrouwbaarheid transcriptie hoog / gemiddeld / laag]": "medium",
      "[tijdstippen of toelichting]": "00:42-00:48: Background noise makes the description of the affected station unclear. The caller clearly states error code E204.",

      "[transcriptie]": "Good afternoon, we have had a problem with machine MCH-2024-00158 since this morning. The machine runs for about twenty minutes and then stops with error code E204. We have already restarted the machine twice. It works again for a short time, but the fault keeps returning. The machine is currently stopped and we need it for today's production.",
      "[probleem]": "The machine stops during production after approximately 20 minutes. Error code E204 appears on the control panel. After restarting, the machine works temporarily, but the problem returns.",
      "[alarmcode of exacte tekst]": "E204",
      "[onderdeel of station]": "Not yet confirmed; the relevant part of the recording is unclear.",
      "[datum / tijd / situatie]": "16-09-2026, during morning production, after approximately 20 minutes of operation.",
      "[eenmalig / af en toe / continu]": "intermittent",
      "[symptomen]": "The machine stops after approximately 20 minutes of production and displays error code E204. Restarting temporarily restores operation.",
      "[acties en resultaten]": "The customer restarted the machine twice. Each restart restored operation temporarily, but the same fault returned.",
      "[onderhoud / instellingen / onderdelen / software]": "No recent changes were mentioned during the call. Confirm whether maintenance, settings, parts, or software have changed.",

      "[ontbrekende informatie]": "The exact alarm text, affected station, operating conditions, software version, and recent maintenance history are not yet known.",
      "[vragen]": "Can you send a photo of the complete alarm message? Which station is affected? Does the fault always occur after approximately 20 minutes? Have any settings, parts, or software changed recently?",
      "[oorzaak]": "The root cause is not yet established. An intermittent fault that develops during operation is suspected, pending confirmation of the meaning of E204 for this machine.",
      "[oorzaken]": "Possible alternatives to investigate include a loose connection, an intermittent sensor fault, or a temperature-related issue. None has been confirmed.",
      "[feiten uit melding en kennisbron]": "The caller reports repeated stops after approximately 20 minutes and only temporary recovery after restarting. No verified knowledge source has been found yet, so the meaning of E204 and the suspected causes remain unconfirmed.",
      "[zekerheid analyse hoog / gemiddeld / laag]": "High",

      "[zoekvraag]": "What does error code E204 mean for machine MCH-2024-00158, and which diagnostic steps apply when the machine stops after approximately 20 minutes?",
      "[zoektermen]": "MCH-2024-00158, E204, intermittent stop, 20 minutes, restart",
      "[bronnen]": "No knowledge sources have been searched yet. Planned sources are the machine manual, service information sheets, and previous service reports.",
      "[documentnaam, documentnummer en versie]": "No verified document has been identified yet.",
      "[map of koppeling]": "Not yet available; add the location of the verified source document.",
      "[datum van bron]": "16-08-2025",
      "[relevantie van zoekresultaat hoog / gemiddeld / laag]": "laag",
      "[conceptadvies]": "Switch off the machine safely, check for visible blockages or loose connections, and do not restart the machine if error code E204 immediately returns. Record the conditions under which the error occurs and wait for further instructions from the engineer.",
      "[documentnaam, documentnummer, versie, paragraaf of pagina]": "Source references are pending. Have the engineer verify each advice step against the applicable machine documentation.",
      "[afwijkingen of onzekerheden]": "The affected station is unclear in the recording, the meaning of E204 has not been verified, and no source document is available yet."
    }"""

    document_path = convert_json_to_docx(json_string_input)
    check_supplied_values_appear(document_path, json_string_input)
    check_remaining_placeholders(document_path)

def test_generator_json_file_input_one_missing_placeholder_throws_not_all_placeholders_replaced_exception():
    script_directory = Path(__file__).resolve().parent

    json_file_input = script_directory / "template" / "text-input-example_one_field_missing.json"

    with pytest.raises(UnreplacedPlaceholdersError):
        check_remaining_placeholders(convert_json_to_docx(json_file_input))

def test_generator_json_string_input_one_missing_placeholder_throws_not_all_placeholders_replaced_exception():

    json_string_input = """{
      "[machinenummer]": "MCH-2024-00158",
      "[distributeur]": "Distributor B.V.",
      "[naam monteur]": "Joost Klein",
      "[contactpersoon]": "Peter Brown",
      "[klant]": "Customer Ltd.",
      "[tijdverschil]": "+1",

      "[datum]": "16-09-2026",
      "[engineer]": "John Smith",
      "[goedgekeurd]": "Smith John",
      "[datum / tijd]": "16-09-2026 / 14:30",

      "[telefoonnummer en/of e-mailadres]": "peter.brown@example.com / +31 6 12345678",
      "[taal]": "English",
      "[naam van Repak medewerker die melding aangenomen heeft]": "John Smith",
      "[machine staat stil / productie beperkt / machine in productie]": "machine is down",
      "[audio link]": "https://audio.website.com/9f33nt9gievnrk",
      "[betrouwbaarheid transcriptie hoog / gemiddeld / laag]": "medium",
      "[tijdstippen of toelichting]": "00:42-00:48: Background noise makes the description of the affected station unclear. The caller clearly states error code E204.",

      "[transcriptie]": "Good afternoon, we have had a problem with machine MCH-2024-00158 since this morning. The machine runs for about twenty minutes and then stops with error code E204. We have already restarted the machine twice. It works again for a short time, but the fault keeps returning. The machine is currently stopped and we need it for today's production.",
      "[probleem]": "The machine stops during production after approximately 20 minutes. Error code E204 appears on the control panel. After restarting, the machine works temporarily, but the problem returns.",
      "[alarmcode of exacte tekst]": "E204",
      "[onderdeel of station]": "Not yet confirmed; the relevant part of the recording is unclear.",
      "[datum / tijd / situatie]": "16-09-2026, during morning production, after approximately 20 minutes of operation.",
      "[eenmalig / af en toe / continu]": "intermittent",
      "[symptomen]": "The machine stops after approximately 20 minutes of production and displays error code E204. Restarting temporarily restores operation.",
      "[acties en resultaten]": "The customer restarted the machine twice. Each restart restored operation temporarily, but the same fault returned.",
      "[onderhoud / instellingen / onderdelen / software]": "No recent changes were mentioned during the call. Confirm whether maintenance, settings, parts, or software have changed.",

      "[ontbrekende informatie]": "The exact alarm text, affected station, operating conditions, software version, and recent maintenance history are not yet known.",

      "[oorzaak]": "The root cause is not yet established. An intermittent fault that develops during operation is suspected, pending confirmation of the meaning of E204 for this machine.",
      "[oorzaken]": "Possible alternatives to investigate include a loose connection, an intermittent sensor fault, or a temperature-related issue. None has been confirmed.",
      "[feiten uit melding en kennisbron]": "The caller reports repeated stops after approximately 20 minutes and only temporary recovery after restarting. No verified knowledge source has been found yet, so the meaning of E204 and the suspected causes remain unconfirmed.",
      "[zekerheid analyse hoog / gemiddeld / laag]": "High",

      "[zoekvraag]": "What does error code E204 mean for machine MCH-2024-00158, and which diagnostic steps apply when the machine stops after approximately 20 minutes?",
      "[zoektermen]": "MCH-2024-00158, E204, intermittent stop, 20 minutes, restart",
      "[bronnen]": "No knowledge sources have been searched yet. Planned sources are the machine manual, service information sheets, and previous service reports.",
      "[documentnaam, documentnummer en versie]": "No verified document has been identified yet.",
      "[map of koppeling]": "Not yet available; add the location of the verified source document.",
      "[datum van bron]": "16-08-2025",
      "[relevantie van zoekresultaat hoog / gemiddeld / laag]": "laag",
      "[conceptadvies]": "Switch off the machine safely, check for visible blockages or loose connections, and do not restart the machine if error code E204 immediately returns. Record the conditions under which the error occurs and wait for further instructions from the engineer.",
      "[documentnaam, documentnummer, versie, paragraaf of pagina]": "Source references are pending. Have the engineer verify each advice step against the applicable machine documentation.",
      "[afwijkingen of onzekerheden]": "The affected station is unclear in the recording, the meaning of E204 has not been verified, and no source document is available yet."
    }"""

    with pytest.raises(UnreplacedPlaceholdersError):
        check_remaining_placeholders(convert_json_to_docx(json_string_input))

def test_generator_json_file_input_one_extra_placeholder_():
    script_directory = Path(__file__).resolve().parent

    json_file_input = script_directory / "template" / "text-input-example_one_field_extra.json"

    with pytest.raises(AssertionError):
        check_supplied_values_appear(convert_json_to_docx(json_file_input), json_file_input)

def test_generator_json_string_input_one_extra_placeholder_():
    json_string_input = """{
      "[machinenummer]": "MCH-2024-00158",
      "[distributeur]": "Distributor B.V.",
      "[naam monteur]": "Joost Klein",
      "[contactpersoon]": "Peter Brown",
      "[klant]": "Customer Ltd.",
      "[tijdverschil]": "+1",

      "[datum]": "16-09-2026",
      "[engineer]": "John Smith",
      "[goedgekeurd]": "Smith John",
      "[datum / tijd]": "16-09-2026 / 14:30",

      "[telefoonnummer en/of e-mailadres]": "peter.brown@example.com / +31 6 12345678",
      "[taal]": "English",
      "[naam van Repak medewerker die melding aangenomen heeft]": "John Smith",
      "[machine staat stil / productie beperkt / machine in productie]": "machine is down",
      "[audio link]": "https://audio.website.com/9f33nt9gievnrk",
      "[betrouwbaarheid transcriptie hoog / gemiddeld / laag]": "medium",
      "[tijdstippen of toelichting]": "00:42-00:48: Background noise makes the description of the affected station unclear. The caller clearly states error code E204.",
      "[Extra placeholder]":"Extra placeholder",

      "[transcriptie]": "Good afternoon, we have had a problem with machine MCH-2024-00158 since this morning. The machine runs for about twenty minutes and then stops with error code E204. We have already restarted the machine twice. It works again for a short time, but the fault keeps returning. The machine is currently stopped and we need it for today's production.",
      "[probleem]": "The machine stops during production after approximately 20 minutes. Error code E204 appears on the control panel. After restarting, the machine works temporarily, but the problem returns.",
      "[alarmcode of exacte tekst]": "E204",
      "[onderdeel of station]": "Not yet confirmed; the relevant part of the recording is unclear.",
      "[datum / tijd / situatie]": "16-09-2026, during morning production, after approximately 20 minutes of operation.",
      "[eenmalig / af en toe / continu]": "intermittent",
      "[symptomen]": "The machine stops after approximately 20 minutes of production and displays error code E204. Restarting temporarily restores operation.",
      "[acties en resultaten]": "The customer restarted the machine twice. Each restart restored operation temporarily, but the same fault returned.",
      "[onderhoud / instellingen / onderdelen / software]": "No recent changes were mentioned during the call. Confirm whether maintenance, settings, parts, or software have changed.",

      "[ontbrekende informatie]": "The exact alarm text, affected station, operating conditions, software version, and recent maintenance history are not yet known.",
      "[vragen]": "Can you send a photo of the complete alarm message? Which station is affected? Does the fault always occur after approximately 20 minutes? Have any settings, parts, or software changed recently?",
      "[oorzaak]": "The root cause is not yet established. An intermittent fault that develops during operation is suspected, pending confirmation of the meaning of E204 for this machine.",
      "[oorzaken]": "Possible alternatives to investigate include a loose connection, an intermittent sensor fault, or a temperature-related issue. None has been confirmed.",
      "[feiten uit melding en kennisbron]": "The caller reports repeated stops after approximately 20 minutes and only temporary recovery after restarting. No verified knowledge source has been found yet, so the meaning of E204 and the suspected causes remain unconfirmed.",
      "[zekerheid analyse hoog / gemiddeld / laag]": "High",

      "[zoekvraag]": "What does error code E204 mean for machine MCH-2024-00158, and which diagnostic steps apply when the machine stops after approximately 20 minutes?",
      "[zoektermen]": "MCH-2024-00158, E204, intermittent stop, 20 minutes, restart",
      "[bronnen]": "No knowledge sources have been searched yet. Planned sources are the machine manual, service information sheets, and previous service reports.",
      "[documentnaam, documentnummer en versie]": "No verified document has been identified yet.",
      "[map of koppeling]": "Not yet available; add the location of the verified source document.",
      "[datum van bron]": "16-08-2025",
      "[relevantie van zoekresultaat hoog / gemiddeld / laag]": "laag",
      "[conceptadvies]": "Switch off the machine safely, check for visible blockages or loose connections, and do not restart the machine if error code E204 immediately returns. Record the conditions under which the error occurs and wait for further instructions from the engineer.",
      "[documentnaam, documentnummer, versie, paragraaf of pagina]": "Source references are pending. Have the engineer verify each advice step against the applicable machine documentation.",
      "[afwijkingen of onzekerheden]": "The affected station is unclear in the recording, the meaning of E204 has not been verified, and no source document is available yet."
    }"""

    with pytest.raises(AssertionError):
        check_supplied_values_appear(convert_json_to_docx(json_string_input), json_string_input)

